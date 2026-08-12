#!/usr/bin/env python3
"""Run the frozen three-arm public-resource public-resource evaluation."""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import replace
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import sys
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from provfold.compare import METHOD_DEFINITIONS, METHOD_DISPLAY_NAMES, compare_methods  # noqa: E402
from provfold.config import AggregationConfig  # noqa: E402
from provfold.core import aggregate, prepare_records  # noqa: E402
from provfold.sensitivity import leave_one_family_out, threshold_surface  # noqa: E402


INPUT_DIR = Path(os.environ.get("PROVFOLD_PUBLIC_INPUT_DIR", Path(__file__).resolve().parent / "source_inputs")).resolve()
RAW_XLSX = INPUT_DIR / "gutMDisorder_v3_Rawdata-based_Disorder_Health.xlsx"
LITERATURE_XLSX = INPUT_DIR / "gutMDisorder_v3_Literature-based_Disorder_Health.xlsx"
TAXONOMY = INPUT_DIR / "ncbi_taxonomy_mapping.csv"
DISBIOME_DIR = INPUT_DIR
CONTRACT = Path(__file__).resolve().parent / "execution_contract.json"
OUTPUT = Path(os.environ.get("PROVFOLD_PUBLIC_OUTPUT", Path(__file__).resolve().parent / "generated_run")).resolve()
GENERATED = OUTPUT / "generated"

ALLOWED_SAMPLE_SOURCES = {"stool", "stool samples", "faeces", "feces"}
DIRECTION_MAP = {
    "increase": "increased",
    "increased": "increased",
    "decrease": "decreased",
    "decreased": "decreased",
}
LOWER_RANKS = {"species", "subspecies", "strain"}
HIGHER_RANKS = {"domain", "kingdom", "phylum", "class", "order", "family"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normal(value: Any) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value).strip())


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, fields: list[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def fields_or(rows: list[dict[str, object]], fallback: list[str]) -> list[str]:
    """Return stable fields even when a scientifically valid output has no rows."""
    return list(rows[0]) if rows else fallback


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def sheet_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    values = sheet.iter_rows(values_only=True)
    header = [normal(value) for value in next(values)]
    rows = [dict(zip(header, row)) for row in values if any(normal(value) for value in row)]
    workbook.close()
    return rows


def strict_indices(
    main_rows: list[dict[str, Any]],
    metadata_rows: list[dict[str, Any]],
    associations: list[dict[str, Any]],
    association_index_field: str,
    allowed_disease_ids: set[str],
) -> set[str]:
    sources: dict[str, set[str]] = defaultdict(set)
    for row in metadata_rows:
        sources[normal(row["Index"])].add(normal(row["SampleSource"]).lower())
    genus_direction_indices = {
        normal(row[association_index_field])
        for row in associations
        if normal(row["Classification"]).lower() == "genus"
        and normal(row["Alteration"]).lower() in DIRECTION_MAP
    }
    return {
        normal(row["Index"])
        for row in main_rows
        if normal(row["HumanOrMouse"]).lower() == "human"
        and normal(row["ResearchType"]).lower() == "gut microbiota associated with phenotype"
        and normal(row["Condition2ID"]) == "D006262"
        and normal(row["Condition1ID"]) in allowed_disease_ids
        and bool(sources.get(normal(row["Index"])))
        and sources[normal(row["Index"])].issubset(ALLOWED_SAMPLE_SOURCES)
        and normal(row["Index"]) in genus_direction_indices
    }


def map_taxon(
    source_rank: str,
    taxid: str,
    taxonomy: dict[str, dict[str, str]],
) -> dict[str, str]:
    resolved = taxonomy.get(taxid)
    if not taxid.isdigit() or resolved is None:
        return {
            "harmonized_taxon": "",
            "harmonized_rank": "",
            "target_taxon": "",
            "mapping_state": "missing_or_unresolved_taxid",
            "lineage_component": "",
            "eligibility": "taxonomy_abstention",
        }
    returned_rank = resolved["returned_rank"]
    if source_rank == "genus" and returned_rank == "genus" and resolved["genus_name"]:
        return {
            "harmonized_taxon": resolved["scientific_name"],
            "harmonized_rank": "genus",
            "target_taxon": resolved["genus_name"],
            "mapping_state": "direct_genus_ncbi",
            "lineage_component": f"ncbi_genus:{resolved['genus_taxid']}",
            "eligibility": "eligible",
        }
    if source_rank in LOWER_RANKS and resolved["genus_name"]:
        harmonized_rank = source_rank if source_rank in {"species", "subspecies"} else "subspecies"
        return {
            "harmonized_taxon": resolved["scientific_name"],
            "harmonized_rank": harmonized_rank,
            "target_taxon": resolved["genus_name"],
            "mapping_state": "folded_lower_rank_ncbi",
            "lineage_component": f"ncbi_genus:{resolved['genus_taxid']}",
            "eligibility": "eligible",
        }
    if source_rank in HIGHER_RANKS:
        reason = "higher_rank_no_downward_mapping"
    elif source_rank in {"genus", *LOWER_RANKS}:
        reason = "source_taxonomy_rank_mismatch"
    else:
        reason = "invalid_or_missing_source_rank"
    return {
        "harmonized_taxon": resolved["scientific_name"],
        "harmonized_rank": source_rank if source_rank in HIGHER_RANKS else "",
        "target_taxon": "",
        "mapping_state": reason,
        "lineage_component": "",
        "eligibility": "taxonomy_abstention",
    }


def build_records(contract: dict[str, Any]) -> tuple[list[dict[str, object]], dict[str, dict[str, str]]]:
    taxonomy = {row["query_taxid"]: row for row in read_csv(TAXONOMY)}
    output: list[dict[str, object]] = []
    labels: dict[str, dict[str, str]] = {"rawdata": {}, "literature": {}}

    raw_main = sheet_rows(RAW_XLSX, "projectr")
    raw_meta = sheet_rows(RAW_XLSX, "metadatar")
    raw_associations = sheet_rows(RAW_XLSX, "associationr")
    raw_ids = set(contract["arms"]["gutmdisorder_rawdata"]["eligible_disease_ids"])
    raw_indices = strict_indices(raw_main, raw_meta, raw_associations, "index", raw_ids)
    raw_main_by_index = {normal(row["Index"]): row for row in raw_main if normal(row["Index"]) in raw_indices}
    for sequence, row in enumerate(raw_associations, start=1):
        index = normal(row["index"])
        if index not in raw_indices or normal(row["Alteration"]).lower() not in DIRECTION_MAP:
            continue
        main = raw_main_by_index[index]
        disease_id = normal(main["Condition1ID"])
        disease_label = normal(main["Condition1"]).replace("，", ", ")
        labels["rawdata"][disease_id] = disease_label
        source_rank = normal(row["Classification"]).lower()
        taxid = normal(row["GutMicrobeNCBIID"])
        mapped = map_taxon(source_rank, taxid, taxonomy)
        project = normal(main["Projectnumber"])
        output.append({
            "record_id": f"GMR-{sequence:05d}",
            "report_id": f"PROJECT:{project}",
            "family_id": f"PROJECT:{project}",
            "family_state": "project_family",
            "comparison_id": f"raw:{index}",
            "evidence_tier": "rawdata_project",
            "source_taxon": normal(row["GutMicrobe"]),
            "source_rank": source_rank or "missing",
            **mapped,
            "direction": DIRECTION_MAP[normal(row["Alteration"]).lower()],
            "source_location": f"gutMDisorder v3 rawdata associationr index={index}; taxid={taxid or 'missing'}",
            "arm": "rawdata",
            "disease_id": disease_id,
            "disease_label": disease_label,
            "pool_status": "quantitative",
            "declared_unit_semantics": "BioProject_resolved_project_family",
        })

    literature_main = sheet_rows(LITERATURE_XLSX, "Literature")
    literature_meta = sheet_rows(LITERATURE_XLSX, "Metadata")
    literature_associations = sheet_rows(LITERATURE_XLSX, "Association")
    literature_ids = set(contract["arms"]["gutmdisorder_literature"]["eligible_disease_ids"])
    allowed_ids = literature_ids | {"", "*"}
    literature_indices = strict_indices(
        literature_main,
        literature_meta,
        literature_associations,
        "Index",
        allowed_ids,
    )
    literature_main_by_index = {normal(row["Index"]): row for row in literature_main if normal(row["Index"]) in literature_indices}
    for sequence, row in enumerate(literature_associations, start=1):
        index = normal(row["Index"])
        if index not in literature_indices or normal(row["Alteration"]).lower() not in DIRECTION_MAP:
            continue
        main = literature_main_by_index[index]
        source_disease_id = normal(main["Condition1ID"])
        disease_label = normal(main["Condition1"]).replace("，", ", ")
        quantitative = source_disease_id in literature_ids
        disease_id = source_disease_id if quantitative else f"UNRESOLVED:{disease_label}"
        labels["literature"][disease_id] = disease_label
        source_rank = normal(row["Classification"]).lower()
        taxid = normal(row["GutMicrobiataNCBIID"])
        mapped = map_taxon(source_rank, taxid, taxonomy)
        pmid = normal(main["PMID"])
        output.append({
            "record_id": f"GML-{sequence:05d}",
            "report_id": f"PMID:{pmid}",
            "family_id": f"PMID:{pmid}",
            "family_state": "publication_family",
            "comparison_id": f"literature:{index}",
            "evidence_tier": "literature_publication",
            "source_taxon": normal(row["GutMicrobe"]),
            "source_rank": source_rank or "missing",
            **mapped,
            "direction": DIRECTION_MAP[normal(row["Alteration"]).lower()],
            "source_location": f"gutMDisorder v3 literature Association index={index}; taxid={taxid or 'missing'}",
            "arm": "literature",
            "disease_id": disease_id,
            "disease_label": disease_label,
            "pool_status": "quantitative" if quantitative else "unresolved_disease_identifier_abstention",
            "declared_unit_semantics": "publication_family_not_cohort_family",
        })
    return output, labels


def config_for(arm: str, records: list[dict[str, object]]) -> AggregationConfig:
    return AggregationConfig.from_dict({
        "analysis_id": f"{arm}_{records[0]['disease_id']}",
        "primary_comparisons": sorted({str(row["comparison_id"]) for row in records}),
        "primary_evidence_tiers": ["rawdata_project" if arm == "rawdata" else "literature_publication"],
        "direction_states": ["increased", "decreased"],
        "eligible_family_states": ["project_family" if arm == "rawdata" else "publication_family"],
        "eligible_mapping_states": ["direct_genus_ncbi", "folded_lower_rank_ncbi"],
        "eligible_record_states": ["eligible"],
        "target_rank": "genus",
        "minimum_support": 2,
        "maximum_opposition": 0,
        "conflict_policy": "abstain",
        "comparator_scope": "primary",
    })


def pairs(rows: Iterable[dict[str, object]], taxon_field: str) -> list[str]:
    return sorted(
        f"{row[taxon_field]}|{row['selected_direction']}"
        for row in rows
        if row.get("selected_direction")
    )


def jaccard(first: set[str], second: set[str]) -> float | None:
    union = first | second
    return None if not union else len(first & second) / len(union)


def augment(rows: Iterable[dict[str, object]], arm: str, disease_id: str, disease_label: str, mode: str = "full") -> list[dict[str, object]]:
    return [{"arm": arm, "disease_id": disease_id, "disease_label": disease_label, "mode": mode, **row} for row in rows]


def disbiome_summary() -> dict[str, object]:
    experiments = json.loads((DISBIOME_DIR / "experiment.json").read_text(encoding="utf-8"))
    accessions = json.loads((DISBIOME_DIR / "publication_accession_numbers.json").read_text(encoding="utf-8"))
    relevant = [
        row for row in experiments
        if normal(row.get("host_type")).lower() == "human"
        and normal(row.get("sample_name")).lower() in {"faeces", "feces", "stool"}
        and normal(row.get("control_name")).lower() == "healthy control"
        and normal(row.get("qualitative_outcome")).lower() in {"elevated", "reduced"}
    ]
    publications = {int(row["publication_id"]) for row in relevant}
    publications_with_accessions = {int(row["publication_id"]) for row in accessions if row.get("accessions")}
    linked = publications & publications_with_accessions
    return {
        "arm": "disbiome",
        "strict_association_rows": len(relevant),
        "strict_publications": len(publications),
        "strict_diseases": len({normal(row.get("disease_name")) for row in relevant}),
        "strict_taxa": len({normal(row.get("organism_name")) for row in relevant}),
        "unique_experiment_identifiers": len({int(row["experiment_id"]) for row in relevant}),
        "experiment_identifier_reuse_count": len(relevant) - len({int(row["experiment_id"]) for row in relevant}),
        "experiment_level_explicit_family_rows": 0,
        "experiment_level_family_determinability_rate": 0.0,
        "quantitative_recurrence_rows": 0,
        "family_indeterminate_abstention_rows": len(relevant),
        "family_indeterminate_abstention_rate": 1.0,
        "publications_with_any_accession": len(linked),
        "publication_accession_presence_rate": len(linked) / len(publications),
        "interpretation": "experiment_id_is_unique_per_association_row_and_does_not_define_a_repeated_experiment_family; publication_id_is_evaluated_separately_as_a_declared_proxy",
    }


def disbiome_publication_family_sensitivity() -> tuple[dict[str, object], list[dict[str, object]], list[dict[str, object]]]:
    """Evaluate the strict Disbiome subset under an explicit publication-family proxy."""

    experiments = json.loads((DISBIOME_DIR / "experiment.json").read_text(encoding="utf-8"))
    relevant = [
        row for row in experiments
        if normal(row.get("host_type")).lower() == "human"
        and normal(row.get("sample_name")).lower() in {"faeces", "feces", "stool"}
        and normal(row.get("control_name")).lower() == "healthy control"
        and normal(row.get("qualitative_outcome")).lower() in {"elevated", "reduced"}
    ]
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in relevant:
        grouped[str(row["disease_id"])].append(row)

    disease_rows: list[dict[str, object]] = []
    selected_rows: list[dict[str, object]] = []
    for disease_id, disease_source_rows in sorted(grouped.items(), key=lambda item: int(item[0])):
        disease_label = normal(disease_source_rows[0].get("disease_name"))
        adapted = []
        for row in disease_source_rows:
            organism_key = f"DBORG:{int(row['organism_id'])}"
            adapted.append({
                "record_id": f"DBEXP:{int(row['experiment_id'])}",
                "report_id": f"DBPUB:{int(row['publication_id'])}",
                "family_id": f"DBPUB:{int(row['publication_id'])}",
                "family_state": "publication_family",
                "comparison_id": f"DBDISEASE:{disease_id}:vs_healthy",
                "evidence_tier": "strict_disbiome_publication_proxy",
                "source_taxon": organism_key,
                "source_rank": "database_entity",
                "harmonized_taxon": organism_key,
                "harmonized_rank": "database_entity",
                "target_taxon": organism_key,
                "mapping_state": "exact_database_entity",
                "direction": "increased" if normal(row["qualitative_outcome"]).lower() == "elevated" else "decreased",
                "lineage_component": organism_key,
                "eligibility": "eligible",
                "source_location": f"Disbiome experiment_id={int(row['experiment_id'])}",
            })
        family_count = len({row["family_id"] for row in adapted})
        status = "evaluable" if family_count >= 3 else "non_evaluable_below_three_publications"
        selected: list[str] = []
        entity_count = 0
        if status == "evaluable":
            config = AggregationConfig.from_dict({
                "analysis_id": f"disbiome_publication_proxy_{disease_id}",
                "primary_comparisons": [f"DBDISEASE:{disease_id}:vs_healthy"],
                "primary_evidence_tiers": ["strict_disbiome_publication_proxy"],
                "direction_states": ["increased", "decreased"],
                "eligible_family_states": ["publication_family"],
                "eligible_mapping_states": ["exact_database_entity"],
                "eligible_record_states": ["eligible"],
                "target_rank": "database_entity",
                "rank_order": ["database_entity"],
                "minimum_support": 2,
                "maximum_opposition": 0,
                "conflict_policy": "abstain",
                "comparator_scope": "primary",
            })
            result = aggregate(adapted, config)
            entity_count = len(result.taxon_summaries)
            selected = sorted(
                f"{row['target_taxon']}|{row['selected_direction']}"
                for row in result.taxon_summaries
                if row["selected_direction"]
            )
            selected_rows.extend({
                "disease_id": disease_id,
                "disease_label": disease_label,
                "organism_entity": row["target_taxon"],
                "selected_direction": row["selected_direction"],
                "publication_family_count": row["family_count"],
                "direction_counts_json": row["direction_counts_json"],
            } for row in result.taxon_summaries if row["selected_direction"])
        disease_rows.append({
            "disease_id": disease_id,
            "disease_label": disease_label,
            "strict_association_row_count": len(disease_source_rows),
            "publication_family_count": family_count,
            "database_organism_entity_count": entity_count,
            "evaluation_status": status,
            "selected_pair_count": len(selected),
            "selected_pairs": ";".join(selected),
            "unit_semantics": "publication_id_as_declared_proxy_not_verified_cohort_family",
            "taxon_semantics": "stable_Disbiome_organism_id_without_taxonomic_rank_inference",
        })
    summary = {
        "strict_association_row_count": len(relevant),
        "disease_count": len(disease_rows),
        "evaluable_disease_count": sum(row["evaluation_status"] == "evaluable" for row in disease_rows),
        "non_evaluable_disease_count": sum(row["evaluation_status"] != "evaluable" for row in disease_rows),
        "publication_family_count": len({int(row["publication_id"]) for row in relevant}),
        "selected_pair_count_across_non_pooled_diseases": sum(int(row["selected_pair_count"]) for row in disease_rows),
        "interpretation": "bounded_publication_family_sensitivity_not_experiment_or_cohort_independence",
    }
    return summary, disease_rows, selected_rows


def main() -> None:
    contract = json.loads(CONTRACT.read_text(encoding="utf-8"))
    identities = {
        "rawdata_xlsx_sha256": RAW_XLSX,
        "literature_xlsx_sha256": LITERATURE_XLSX,
        "taxonomy_mapping_sha256": TAXONOMY,
        "disbiome_experiment_sha256": DISBIOME_DIR / "experiment.json",
        "disbiome_accession_sha256": DISBIOME_DIR / "publication_accession_numbers.json",
    }
    for key, path in identities.items():
        if sha256(path) != contract["source_identities"][key]:
            raise SystemExit(f"frozen public-resource source identity mismatch: {key}")

    records, labels = build_records(contract)
    quantitative = [row for row in records if row["pool_status"] == "quantitative"]
    unresolved = [row for row in records if row["pool_status"] != "quantitative"]

    raw_quantitative = [row for row in quantitative if row["arm"] == "rawdata"]
    literature_quantitative = [row for row in quantitative if row["arm"] == "literature"]
    observed_contract_counts = {
        "raw_comparisons": len({str(row["comparison_id"]) for row in raw_quantitative}),
        "raw_projects": len({str(row["family_id"]) for row in raw_quantitative}),
        "literature_comparisons": len({str(row["comparison_id"]) for row in literature_quantitative}),
        "literature_publications": len({str(row["family_id"]) for row in literature_quantitative}),
        "unresolved_comparisons": len({str(row["comparison_id"]) for row in unresolved}),
        "unresolved_publications": len({str(row["family_id"]) for row in unresolved}),
        "unresolved_genus_rows": sum(str(row["source_rank"]) == "genus" for row in unresolved),
    }
    expected_contract_counts = {
        "raw_comparisons": contract["arms"]["gutmdisorder_rawdata"]["comparison_count"],
        "raw_projects": contract["arms"]["gutmdisorder_rawdata"]["distinct_project_count_across_diseases"],
        "literature_comparisons": contract["arms"]["gutmdisorder_literature"]["quantitative_comparison_count"],
        "literature_publications": contract["arms"]["gutmdisorder_literature"]["distinct_publication_count"],
        "unresolved_comparisons": contract["arms"]["gutmdisorder_literature"]["unresolved_identifier_comparison_count"],
        "unresolved_publications": contract["arms"]["gutmdisorder_literature"]["unresolved_identifier_publication_count"],
        "unresolved_genus_rows": contract["arms"]["gutmdisorder_literature"]["unresolved_identifier_genus_direction_row_count"],
    }
    if observed_contract_counts != expected_contract_counts:
        raise SystemExit(
            "frozen public-resource corpus count mismatch: "
            f"observed={observed_contract_counts}; expected={expected_contract_counts}"
        )
    GENERATED.mkdir(parents=True, exist_ok=True)

    record_fields = list(records[0])
    write_csv(GENERATED / "adapted_records.csv", record_fields, records)

    arm_disease_summary: list[dict[str, object]] = []
    method_summary: list[dict[str, object]] = []
    rank_summary: list[dict[str, object]] = []
    omission_summary: list[dict[str, object]] = []
    detailed_evidence_units: list[dict[str, object]] = []
    detailed_family_states: list[dict[str, object]] = []
    detailed_taxon_summaries: list[dict[str, object]] = []
    detailed_abstentions: list[dict[str, object]] = []
    detailed_methods: list[dict[str, object]] = []
    detailed_surface: list[dict[str, object]] = []
    detailed_omission: list[dict[str, object]] = []
    detailed_direct_surface: list[dict[str, object]] = []
    detailed_rank_contrast: list[dict[str, object]] = []
    results_by_key: dict[tuple[str, str], dict[str, object]] = {}

    grouped: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in quantitative:
        grouped[(str(row["arm"]), str(row["disease_id"]))].append(row)

    for (arm, disease_id), disease_records in sorted(grouped.items()):
        label = labels[arm][disease_id]
        config = config_for(arm, disease_records)
        voting, abstentions, validation = prepare_records(disease_records, config)
        evaluable_families = sorted({str(row["family_id"]) for row in voting})
        status = "evaluable" if len(evaluable_families) >= 3 else "non_evaluable_below_three_units"
        base_summary = {
            "arm": arm,
            "disease_id": disease_id,
            "disease_label": label,
            "declared_comparison_count": len({str(row["comparison_id"]) for row in disease_records}),
            "declared_family_count": len({str(row["family_id"]) for row in disease_records}),
            "directional_source_row_count": len(disease_records),
            "voting_row_count": len(voting),
            "abstention_row_count": len(abstentions),
            "evaluable_family_count": len(evaluable_families),
            "validation_warning_count": sum(item["level"] == "warning" for item in validation.issues),
            "evaluation_status": status,
        }
        if status != "evaluable":
            arm_disease_summary.append({
                **base_summary,
                "target_genus_count": 0,
                "primary_selected_pair_count": 0,
                "primary_selected_pairs": "",
                "threshold_setting_count": 0,
            })
            continue

        full = aggregate(disease_records, config)
        comparison = compare_methods(disease_records, config)
        surface = threshold_surface(disease_records, config)
        omission = leave_one_family_out(disease_records, config)
        reference_pairs = pairs(full.taxon_summaries, "target_taxon")
        arm_disease_summary.append({
            **base_summary,
            "target_genus_count": len(full.taxon_summaries),
            "primary_selected_pair_count": len(reference_pairs),
            "primary_selected_pairs": ";".join(reference_pairs),
            "threshold_setting_count": len(evaluable_families) ** 2,
        })

        for method in sorted({str(row["method"]) for row in comparison}):
            method_rows = [row for row in comparison if row["method"] == method]
            selected = pairs(method_rows, "taxon_key")
            method_summary.append({
                "arm": arm,
                "disease_id": disease_id,
                "disease_label": label,
                "method": method,
                "method_display_name": METHOD_DISPLAY_NAMES[method],
                "unit_semantics": METHOD_DEFINITIONS[method],
                "taxon_summary_count": len(method_rows),
                "vote_count_sum": sum(int(row["vote_count"]) for row in method_rows),
                "selected_pair_count": len(selected),
                "selected_pairs": ";".join(selected),
            })

        direct_records = [row for row in disease_records if row["mapping_state"] == "direct_genus_ncbi"]
        direct_voting, _, _ = prepare_records(direct_records, config)
        direct_families = sorted({str(row["family_id"]) for row in direct_voting})
        direct_full = aggregate(direct_records, config) if direct_records else None
        direct_pairs = pairs(direct_full.taxon_summaries, "target_taxon") if direct_full else []
        rank_summary.append({
            "arm": arm,
            "disease_id": disease_id,
            "disease_label": label,
            "full_voting_row_count": len(voting),
            "direct_genus_voting_row_count": len(direct_voting),
            "full_evidence_unit_count": len(full.evidence_units),
            "direct_genus_evidence_unit_count": len(direct_full.evidence_units) if direct_full else 0,
            "full_family_taxon_state_count": len(full.family_taxon_states),
            "direct_genus_family_taxon_state_count": len(direct_full.family_taxon_states) if direct_full else 0,
            "full_selected_pair_count": len(reference_pairs),
            "direct_genus_selected_pair_count": len(direct_pairs),
            "added_by_lower_rank_folding": ";".join(sorted(set(reference_pairs) - set(direct_pairs))),
            "removed_after_lower_rank_folding": ";".join(sorted(set(direct_pairs) - set(reference_pairs))),
            "primary_pair_jaccard": "" if jaccard(set(reference_pairs), set(direct_pairs)) is None else round(jaccard(set(reference_pairs), set(direct_pairs)), 8),
        })

        direct_surface: list[dict[str, object]] = []
        if direct_families:
            direct_surface = threshold_surface(direct_records, config)
            full_by_setting: dict[tuple[int, int], list[dict[str, object]]] = defaultdict(list)
            direct_by_setting: dict[tuple[int, int], list[dict[str, object]]] = defaultdict(list)
            for row in surface:
                full_by_setting[(int(row["minimum_support"]), int(row["maximum_opposition"]))].append(row)
            for row in direct_surface:
                direct_by_setting[(int(row["minimum_support"]), int(row["maximum_opposition"]))].append(row)
            for setting in sorted(set(full_by_setting) & set(direct_by_setting)):
                first = set(pairs(full_by_setting[setting], "target_taxon"))
                second = set(pairs(direct_by_setting[setting], "target_taxon"))
                detailed_rank_contrast.append({
                    "arm": arm,
                    "disease_id": disease_id,
                    "disease_label": label,
                    "minimum_support": setting[0],
                    "maximum_opposition": setting[1],
                    "full_selected_pair_count": len(first),
                    "direct_genus_selected_pair_count": len(second),
                    "added_by_lower_rank_folding": ";".join(sorted(first - second)),
                    "removed_after_lower_rank_folding": ";".join(sorted(second - first)),
                    "Jaccard": "" if jaccard(first, second) is None else round(jaccard(first, second), 8),
                })

        omission_by_family: dict[str, list[dict[str, object]]] = defaultdict(list)
        for row in omission:
            omission_by_family[str(row["omitted_family_id"])].append(row)
        evaluable_omission_sets: list[set[str]] = []
        non_evaluable_omissions = 0
        for family_id, rows in sorted(omission_by_family.items()):
            remaining = len(evaluable_families) - 1
            omission_status = "evaluable" if remaining >= 3 else "non_evaluable_after_omission"
            selected = set(pairs(rows, "target_taxon"))
            if omission_status == "evaluable":
                evaluable_omission_sets.append(selected)
            else:
                non_evaluable_omissions += 1
            for row in rows:
                detailed_omission.append({
                    "arm": arm,
                    "disease_id": disease_id,
                    "disease_label": label,
                    "remaining_evaluable_family_count": remaining,
                    "omission_evaluation_status": omission_status,
                    **row,
                })
        surviving = set(reference_pairs)
        for selected in evaluable_omission_sets:
            surviving &= selected
        omission_summary.append({
            "arm": arm,
            "disease_id": disease_id,
            "disease_label": label,
            "reference_selected_pair_count": len(reference_pairs),
            "omission_count": len(omission_by_family),
            "evaluable_omission_count": len(evaluable_omission_sets),
            "non_evaluable_omission_count": non_evaluable_omissions,
            "survives_all_evaluable_omissions_count": len(surviving) if evaluable_omission_sets else "",
            "survives_all_evaluable_omissions_pairs": ";".join(sorted(surviving)) if evaluable_omission_sets else "",
            "interpretation": "assessable" if evaluable_omission_sets else "not_assessable_all_omissions_below_three_units",
        })

        detailed_evidence_units.extend(augment(full.evidence_units, arm, disease_id, label))
        detailed_family_states.extend(augment(full.family_taxon_states, arm, disease_id, label))
        detailed_taxon_summaries.extend(augment(full.taxon_summaries, arm, disease_id, label))
        detailed_abstentions.extend(augment(full.abstentions, arm, disease_id, label))
        detailed_methods.extend(augment(comparison, arm, disease_id, label))
        detailed_surface.extend(augment(surface, arm, disease_id, label))
        detailed_direct_surface.extend(augment(direct_surface, arm, disease_id, label, "direct_genus_only"))
        results_by_key[(arm, disease_id)] = {
            "family_states": full.family_taxon_states,
            "surface": surface,
            "evaluable_family_count": len(evaluable_families),
        }

    overlap = contract["cross_arm_concordance"]["overlap_disease_ids"]
    balance_rows: list[dict[str, object]] = []
    concordance_summary: list[dict[str, object]] = []
    cross_jaccard: list[dict[str, object]] = []
    for disease_id in overlap:
        raw_result = results_by_key.get(("rawdata", disease_id))
        literature_result = results_by_key.get(("literature", disease_id))
        if raw_result is None or literature_result is None:
            concordance_summary.append({
                "disease_id": disease_id,
                "disease_label": labels["rawdata"].get(disease_id, labels["literature"].get(disease_id, "")),
                "evaluation_status": "non_evaluable_in_at_least_one_arm",
                "shared_genus_count": 0,
                "nonzero_balance_comparison_count": 0,
                "sign_agreement_count": 0,
                "sign_agreement_rate": "",
            })
            continue

        def balances(states: list[dict[str, object]]) -> dict[str, dict[str, object]]:
            grouped_states: dict[str, list[dict[str, object]]] = defaultdict(list)
            for row in states:
                grouped_states[str(row["target_taxon"])].append(row)
            result: dict[str, dict[str, object]] = {}
            for taxon, rows in grouped_states.items():
                increased = sum(str(row["family_direction_state"]) == "increased" for row in rows)
                decreased = sum(str(row["family_direction_state"]) == "decreased" for row in rows)
                conflicts = sum(str(row["family_direction_state"]) == "conflict" for row in rows)
                denominator = increased + decreased
                result[taxon] = {
                    "increased": increased,
                    "decreased": decreased,
                    "conflicts": conflicts,
                    "balance": None if denominator == 0 else (increased - decreased) / denominator,
                }
            return result

        raw_balance = balances(raw_result["family_states"])
        literature_balance = balances(literature_result["family_states"])
        shared = sorted(set(raw_balance) & set(literature_balance))
        agreement_denominator = 0
        agreement_count = 0
        for genus in shared:
            first = raw_balance[genus]
            second = literature_balance[genus]
            first_sign = 0 if first["balance"] == 0 else (1 if first["balance"] and first["balance"] > 0 else -1)
            second_sign = 0 if second["balance"] == 0 else (1 if second["balance"] and second["balance"] > 0 else -1)
            agreement = "excluded_zero_balance"
            if first_sign != 0 and second_sign != 0:
                agreement_denominator += 1
                agreement = "agree" if first_sign == second_sign else "disagree"
                agreement_count += first_sign == second_sign
            balance_rows.append({
                "disease_id": disease_id,
                "disease_label": labels["rawdata"].get(disease_id, labels["literature"].get(disease_id, "")),
                "genus": genus,
                "rawdata_increased_units": first["increased"],
                "rawdata_decreased_units": first["decreased"],
                "rawdata_conflict_units": first["conflicts"],
                "rawdata_direction_balance": round(first["balance"], 8) if first["balance"] is not None else "",
                "literature_increased_units": second["increased"],
                "literature_decreased_units": second["decreased"],
                "literature_conflict_units": second["conflicts"],
                "literature_direction_balance": round(second["balance"], 8) if second["balance"] is not None else "",
                "sign_agreement_status": agreement,
            })
        concordance_summary.append({
            "disease_id": disease_id,
            "disease_label": labels["rawdata"].get(disease_id, labels["literature"].get(disease_id, "")),
            "evaluation_status": "descriptive_non_pooled",
            "shared_genus_count": len(shared),
            "nonzero_balance_comparison_count": agreement_denominator,
            "sign_agreement_count": agreement_count,
            "sign_agreement_rate": "" if agreement_denominator == 0 else round(agreement_count / agreement_denominator, 8),
        })

        def surface_sets(rows: list[dict[str, object]]) -> dict[tuple[int, int], set[str]]:
            output: dict[tuple[int, int], list[dict[str, object]]] = defaultdict(list)
            for row in rows:
                output[(int(row["minimum_support"]), int(row["maximum_opposition"]))].append(row)
            return {setting: set(pairs(values, "target_taxon")) for setting, values in output.items()}

        raw_surface = surface_sets(raw_result["surface"])
        literature_surface = surface_sets(literature_result["surface"])
        for setting in sorted(set(raw_surface) & set(literature_surface)):
            first = raw_surface[setting]
            second = literature_surface[setting]
            cross_jaccard.append({
                "disease_id": disease_id,
                "disease_label": labels["rawdata"].get(disease_id, labels["literature"].get(disease_id, "")),
                "minimum_support": setting[0],
                "maximum_opposition": setting[1],
                "rawdata_selected_pair_count": len(first),
                "literature_selected_pair_count": len(second),
                "intersection_count": len(first & second),
                "union_count": len(first | second),
                "Jaccard": "" if jaccard(first, second) is None else round(jaccard(first, second), 8),
            })

    disbiome = disbiome_summary()
    disbiome_publication_proxy, disbiome_disease_rows, disbiome_selected_rows = disbiome_publication_family_sensitivity()
    unresolved_comparisons = len({str(row["comparison_id"]) for row in unresolved})
    unresolved_publications = len({str(row["family_id"]) for row in unresolved})
    unresolved_genus_rows = sum(str(row["source_rank"]) == "genus" for row in unresolved)
    taxonomy_rows = read_csv(TAXONOMY)
    abstention_summary = [
        {
            "arm": arm,
            "scope": "quantitative_pool",
            "directional_source_row_count": sum(row["arm"] == arm for row in quantitative),
            "voting_row_count": sum(row["arm"] == arm and row["eligibility"] == "eligible" for row in quantitative),
            "taxonomy_abstention_row_count": sum(row["arm"] == arm and row["eligibility"] != "eligible" for row in quantitative),
            "disease_identifier_abstention_row_count": 0,
            "family_indeterminate_abstention_row_count": 0,
        }
        for arm in ("rawdata", "literature")
    ]
    abstention_summary.extend([
        {
            "arm": "literature",
            "scope": "unresolved_disease_identifier_audit",
            "directional_source_row_count": len(unresolved),
            "voting_row_count": 0,
            "taxonomy_abstention_row_count": sum(row["eligibility"] != "eligible" for row in unresolved),
            "disease_identifier_abstention_row_count": len(unresolved),
            "family_indeterminate_abstention_row_count": 0,
        },
        {
            "arm": "disbiome",
            "scope": "strict_experiment_family_definition",
            "directional_source_row_count": disbiome["strict_association_rows"],
            "voting_row_count": 0,
            "taxonomy_abstention_row_count": 0,
            "disease_identifier_abstention_row_count": 0,
            "family_indeterminate_abstention_row_count": disbiome["family_indeterminate_abstention_rows"],
        },
        {
            "arm": "disbiome",
            "scope": "publication_family_proxy_sensitivity",
            "directional_source_row_count": disbiome_publication_proxy["strict_association_row_count"],
            "voting_row_count": disbiome_publication_proxy["strict_association_row_count"],
            "taxonomy_abstention_row_count": 0,
            "disease_identifier_abstention_row_count": 0,
            "family_indeterminate_abstention_row_count": 0,
        },
    ])

    write_csv(OUTPUT / "arm_disease_summary.csv", fields_or(arm_disease_summary, ["arm", "disease_id"]), arm_disease_summary)
    write_csv(OUTPUT / "primary_method_summary.csv", fields_or(method_summary, ["arm", "disease_id", "method"]), method_summary)
    write_csv(OUTPUT / "rank_folding_summary.csv", fields_or(rank_summary, ["arm", "disease_id"]), rank_summary)
    write_csv(OUTPUT / "omission_summary.csv", fields_or(omission_summary, ["arm", "disease_id"]), omission_summary)
    write_csv(
        OUTPUT / "cross_arm_direction_balance.csv",
        fields_or(balance_rows, [
            "disease_id", "disease_label", "genus", "rawdata_increased_units",
            "rawdata_decreased_units", "rawdata_conflict_units", "rawdata_direction_balance",
            "literature_increased_units", "literature_decreased_units",
            "literature_conflict_units", "literature_direction_balance", "sign_agreement_status",
        ]),
        balance_rows,
    )
    write_csv(OUTPUT / "cross_arm_concordance_summary.csv", fields_or(concordance_summary, ["disease_id", "evaluation_status"]), concordance_summary)
    write_csv(OUTPUT / "cross_arm_jaccard_surface.csv", fields_or(cross_jaccard, ["disease_id", "minimum_support", "maximum_opposition", "Jaccard"]), cross_jaccard)
    write_csv(OUTPUT / "disbiome_determinability.csv", list(disbiome), [disbiome])
    write_csv(
        OUTPUT / "disbiome_publication_family_summary.csv",
        fields_or(disbiome_disease_rows, ["disease_id", "evaluation_status"]),
        disbiome_disease_rows,
    )
    write_csv(
        OUTPUT / "disbiome_publication_family_selected_pairs.csv",
        fields_or(disbiome_selected_rows, ["disease_id", "organism_entity", "selected_direction"]),
        disbiome_selected_rows,
    )
    write_csv(OUTPUT / "abstention_summary.csv", fields_or(abstention_summary, ["arm", "scope"]), abstention_summary)

    generated_sets = [
        ("evidence_units.csv", detailed_evidence_units),
        ("family_taxon_states.csv", detailed_family_states),
        ("taxon_summaries.csv", detailed_taxon_summaries),
        ("abstentions.csv", detailed_abstentions),
        ("method_comparison.csv", detailed_methods),
        ("threshold_surface.csv", detailed_surface),
        ("leave_one_family_out.csv", detailed_omission),
        ("direct_genus_threshold_surface.csv", detailed_direct_surface),
        ("rank_contrast_surface.csv", detailed_rank_contrast),
    ]
    for name, rows in generated_sets:
        if rows:
            write_csv(GENERATED / name, list(rows[0]), rows)

    result = {
        "schema_version": "1.0",
        "date": "2026-08-12",
        "status": "COMPLETE_PENDING_INDEPENDENT_VALIDATION",
        "contract_sha256": sha256(CONTRACT),
        "quantitative_arm_disease_count": len(arm_disease_summary),
        "evaluable_arm_disease_count": sum(row["evaluation_status"] == "evaluable" for row in arm_disease_summary),
        "non_evaluable_arm_disease_count": sum(row["evaluation_status"] != "evaluable" for row in arm_disease_summary),
        "rawdata_quantitative_disease_count": sum(row["arm"] == "rawdata" for row in arm_disease_summary),
        "literature_quantitative_disease_count": sum(row["arm"] == "literature" for row in arm_disease_summary),
        "unresolved_literature": {
            "comparison_count": unresolved_comparisons,
            "publication_count": unresolved_publications,
            "directional_all_rank_row_count": len(unresolved),
            "genus_direction_row_count": unresolved_genus_rows,
            "quantitative_pool_count": 0,
        },
        "taxonomy": {
            "query_count": len(taxonomy_rows),
            "resolved_count": sum(row["resolution_status"].startswith("resolved") for row in taxonomy_rows),
            "alias_resolved_count": sum(row["resolution_status"] == "resolved_aka_taxid" for row in taxonomy_rows),
            "genus_lineage_available_count": sum(bool(row["genus_taxid"]) for row in taxonomy_rows),
        },
        "disbiome": disbiome,
        "disbiome_publication_family_sensitivity": disbiome_publication_proxy,
        "cross_arm_predeclared_shared_disease_count": len(overlap),
        "cross_arm_evaluable_disease_count": sum(
            row["evaluation_status"] == "descriptive_non_pooled"
            for row in concordance_summary
        ),
        "cross_arm_balance_row_count": len(balance_rows),
        "cross_arm_jaccard_row_count": len(cross_jaccard),
        "arms_pooled": False,
        "publication_families_called_cohort_families": False,
    }
    write_json(OUTPUT / "results.json", result)

    manifest_rows = []
    manifest_exclusions = {
        "artifact_manifest.csv",
        "completion_status.json",
        "independent_validation.json",
        "independent_validation.md",
    }
    for path in sorted(item for item in OUTPUT.rglob("*") if item.is_file() and item.name not in manifest_exclusions):
        manifest_rows.append({
            "relative_path": path.relative_to(OUTPUT).as_posix(),
            "storage_class": "local_ignored_row_level" if GENERATED in path.parents else "tracked_privacy_bounded_summary",
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    write_csv(OUTPUT / "artifact_manifest.csv", list(manifest_rows[0]), manifest_rows)


if __name__ == "__main__":
    main()
